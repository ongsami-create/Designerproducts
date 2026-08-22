import json
from openpyxl import load_workbook

# Load Excel file
wb = load_workbook(r'D:\Users\sami_\OneDrive\Desktop\2026 DESIGNER PRICE LIST 10-4-26.xlsx')
ws = wb.active

products = []

# Skip first 3 rows (empty, header, empty), start from row 4
for idx, row in enumerate(ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True), 1):
    # Extract values (0-indexed)
    area = str(row[0]).strip() if row[0] else ''
    area_zh = str(row[1]).strip() if row[1] else ''
    item_a = str(row[2]).strip() if row[2] else ''
    item_a_zh = str(row[3]).strip() if row[3] else ''
    item_b = str(row[4]).strip() if row[4] else ''
    item_b_zh = str(row[5]).strip() if row[5] else ''
    description = str(row[6]).strip() if row[6] else ''
    description_zh = str(row[7]).strip() if row[7] else ''
    note_a = str(row[8]).strip() if row[8] else ''
    note_a_zh = str(row[9]).strip() if row[9] else ''
    note_b = str(row[10]).strip() if row[10] else ''
    note_b_zh = str(row[11]).strip() if row[11] else ''
    
    # Price handling - try to convert to float
    def parse_price(val):
        if val is None:
            return 0
        if isinstance(val, (int, float)):
            return float(val)
        val_str = str(val).strip()
        if not val_str:
            return 0
        # Remove any formulas like =SUM(...)
        if val_str.startswith('='):
            return 0
        try:
            return float(val_str)
        except:
            return 0
    
    unit_price = parse_price(row[12])
    high_end_price = parse_price(row[13])
    dealer_price = parse_price(row[14])
    cost_price = parse_price(row[15])
    
    input_dimension = str(row[16]).strip() if row[16] else ''
    algorithm = str(row[17]).strip() if row[17] else ''
    algorithm_notes = str(row[19]).strip() if row[19] else ''
    
    # Restrictions
    restriction_d = str(row[20]).strip() if row[20] else ''
    restriction_h = str(row[21]).strip() if row[21] else ''
    restriction_w = str(row[22]).strip() if row[22] else ''
    
    # Special modes
    special_mode_d = str(row[23]).strip() if row[23] else ''
    special_mode_h = str(row[24]).strip() if row[24] else ''
    special_mode_w = str(row[25]).strip() if row[25] else ''
    
    # Section 2 restrictions (exceed limits)
    restriction_d2 = str(row[26]).strip() if row[26] else ''
    restriction_h2 = str(row[27]).strip() if row[27] else ''
    restriction_w2 = str(row[28]).strip() if row[28] else ''
    special_mode_whd = str(row[29]).strip() if row[29] else ''
    
    # Skip rows without area and item_b (product name)
    if not area or not item_b:
        continue
    
    # Create product object
    product = {
        "id": f"p_{idx}",
        "area": area,
        "areaZh": area_zh,
        "itemA": item_a,
        "itemAZh": item_a_zh,
        "itemB": item_b,
        "itemBZh": item_b_zh,
        "description": description,
        "descriptionZh": description_zh,
        "unitPrice": unit_price,
        "highEndPrice": high_end_price,
        "dealerPrice": dealer_price,
        "costPrice": cost_price,
        "noteA": note_a,
        "noteAZh": note_a_zh,
        "noteB": note_b,
        "noteBZh": note_b_zh,
        "inputDimension": input_dimension,
        "algorithm": algorithm,
        "algorithmNotes": algorithm_notes,
        "restrictionD": restriction_d,
        "restrictionH": restriction_h,
        "restrictionW": restriction_w,
        "specialModeD": special_mode_d,
        "specialModeH": special_mode_h,
        "specialModeW": special_mode_w,
        "restrictionD2": restriction_d2,
        "restrictionH2": restriction_h2,
        "restrictionW2": restriction_w2,
        "specialModeWHD": special_mode_whd
    }
    products.append(product)

# Statistics
print(f'Total products: {len(products)}')

# Unique areas
areas = set(p['area'] for p in products)
print(f'Unique areas: {len(areas)}')

# Unique products (itemB)
unique_products = set(p['itemB'] for p in products)
print(f'Unique product names: {len(unique_products)}')

# Save to JSON
output_data = {
    "products": products,
    "meta": {
        "total": len(products),
        "areas": len(areas),
        "updated": "2026-04-21"
    }
}

with open('products.json', 'w', encoding='utf-8') as f:
    json.dump(output_data, f, ensure_ascii=False, indent=2)

print('products.json saved successfully!')
