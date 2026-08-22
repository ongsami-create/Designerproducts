import json
from openpyxl import load_workbook

# Load Excel file
wb = load_workbook(r'D:\Users\sami_\OneDrive\Desktop\2026 DESIGNER PRICE LIST 5-4-26.xlsx')
ws = wb.active

# Data containers
areas_dict = {}
projects_dict = {}
data = {'areas': []}

# Skip first 3 rows (empty, header, empty), start from row 4
for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
    # Column mapping (0-indexed):
    # 0: Area, 2: Project, 4: Product name, 6: Description
    # 8: Note A, 10: Note B, 12: Selling Price, 13: Dealer Price
    # 14: Cost Price, 15: Input Dimension, 16: Algorithm
    # 17: Algorithm Notes, 18: Size Limit
    
    area_name = row[0]
    project_name = row[2]
    product_name = row[4]  # Column E (index 4)
    description = row[6]
    note_a = row[8]
    note_b = row[10]
    selling_price = row[12]
    dealer_price = row[13]
    cost_price = row[14]
    input_dim = row[15]
    algorithm = row[16]
    algo_notes = row[17]
    size_limit = row[18]
    
    area_name_str = str(area_name).strip() if area_name else ''
    project_name_str = str(project_name).strip() if project_name else ''
    product_name_str = str(product_name).strip() if product_name else ''
    
    # Skip rows without area/project
    if not area_name_str or not project_name_str:
        continue
    
    # Create or get Area
    if area_name_str not in areas_dict:
        area_obj = {
            'id': f'area_{len(areas_dict)+1}',
            'name': area_name_str,
            'name_en': '',
            'order': len(areas_dict),
            'projects': []
        }
        areas_dict[area_name_str] = area_obj
        data['areas'].append(area_obj)
        print(f"New Area: {area_name_str}")
    
    current_area = areas_dict[area_name_str]
    
    # Create or get Project within this Area
    project_key = f"{area_name_str}|{project_name_str}"
    if project_key not in projects_dict:
        project_obj = {
            'id': f'project_{len(current_area["projects"])+1}',
            'name': project_name_str,
            'name_en': '',
            'order': len(current_area['projects']),
            'products': []
        }
        projects_dict[project_key] = project_obj
        current_area['projects'].append(project_obj)
        print(f"  New Project: {project_name_str}")
    
    current_project = projects_dict[project_key]
    
    # Handle Product - each row with product name is a product
    if product_name_str:
        # Convert price values
        try:
            sp = float(selling_price) if selling_price else 0
        except:
            sp = 0
        try:
            dp = float(dealer_price) if dealer_price else 0
        except:
            dp = 0
        try:
            cp = float(cost_price) if cost_price else 0
        except:
            cp = 0
            
        product = {
            'id': f'product_{len(current_project["products"])+1}',
            'name': product_name_str,
            'name_en': '',
            'description': str(description).strip() if description else '',
            'note_a': str(note_a).strip() if note_a else '',
            'note_b': str(note_b).strip() if note_b else '',
            'selling_price': sp,
            'cost_price': cp,
            'dealer_price': dp,
            'unit': 'UNIT',
            'input_dimension': str(input_dim).strip() if input_dim else '',
            'algorithm': str(algorithm).upper().strip() if algorithm else 'UNIT',
            'algorithm_notes': str(algo_notes).strip() if algo_notes else '',
            'size_limit': str(size_limit).strip() if size_limit else '',
            'high_end_enabled': False,
            'high_end_percentage': 25,
            'order': len(current_project['products'])
        }
        current_project['products'].append(product)

# Count stats
areas = len(data['areas'])
projects = sum(len(a['projects']) for a in data['areas'])
products = sum(len(p['products']) for a in data['areas'] for p in a['projects'])

print(f'\n=== Final Stats ===')
print(f'Areas: {areas}')
print(f'Projects: {projects}')
print(f'Products: {products}')

# Save JSON
with open('products.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print('\nproducts.json saved successfully!')
